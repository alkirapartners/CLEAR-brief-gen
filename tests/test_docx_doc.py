"""Word export of a brief stored as a JSON document: plain, real styles, safe with untrusted text."""

import copy
import io
import json
import zipfile
from datetime import datetime, timezone

import pytest
from docx import Document

import brief_doc
import docx_doc
from tests import export_readers as readers
from tests.brief_fixtures import SAMPLE_DOC, make_doc

WHEN = datetime(2026, 10, 5, 12, 0)
BULLET, NUMBERED, CONTINUED = "List Bullet", "List Number", "List Continue"


def _render(doc, language=None):
    return docx_doc.render(doc, WHEN, language)


def _styles(data, style):
    return [text for name, text in readers.docx_paragraphs(data) if name == style]


def test_the_headings_are_real_word_headings_in_the_order_of_the_page():
    assert readers.docx_headings(_render(make_doc())) == [
        ("Title", "Northwind Energy"),
        ("Heading 1", "Alkira Fit"),
        ("Heading 1", "Why this account, why now"),
        ("Heading 2", "Angle 1: Hand-built Azure network"),
        ("Heading 2", "Angle 2: Lubricants separation"),
        ("Heading 1", "Ask this"),
        ("Heading 1", "For the engineer"),
        ("Heading 1", "Who to talk to"),
        ("Heading 1", "What we couldn't confirm"),
        ("Heading 1", "What would raise the score"),
        ("Heading 1", "References"),
    ]


def test_every_part_of_the_brief_is_in_the_document():
    text = readers.docx_text(_render(make_doc()))
    for expected in (
        "Alkira opportunity brief", "Researched 5 Oct 2026",
        "Northwind Energy Corporation", "Public (NYSE: NWE)", "northwind.example",
        "Which company: Researched Northwind Energy Corporation",
        "HQ: Dallas, TX", "Revenue: $28B", "Employees: 5,200", "Industry: Refining",
        "5 / 5", "Strong fit", "Strong fit: a hand-built Azure network",
        "Lead with: Open with the Azure hub build", "Cloud and network: Azure, ExpressRoute and Virtual WAN, SD-WAN",
        "Use case: Multi-cloud", "Use case: M&A", "Pending deal, announced 20 Feb 2026",
        "“The separation is expected to be completed over the next 12-18 months.”",
        "What Alkira does: Alkira replaces hand-built hubs",
        "A network engineer posting lists ExpressRoute and a Virtual WAN hub-and-spoke. (source dated 23 Sep 2026) [1]",
        "Customer story: Koch Industries: Significant reduction in network complexity",
        "Who builds a new Virtual WAN hub today, and how long does one take?",
        "Angle 1: Multi-cloud", "Listen for: hand-built hubs, weeks of lead time",
        "Alkira angle: A new region is a design change deployed in a day.",
        "Clouds: Azure is the primary cloud. [1]", "Data centers: Not found in public sources.",
        "Dana Ruiz, Chief Information Officer: Named in the annual report. [2]",
        "Director of Network Engineering: Owns the Azure network. [1]",
        "Who owns the WAN contract.", "A dated SD-WAN or MPLS renewal.",
        "[1] Senior Network Engineer posting (First-hand; source dated 23 Sep 2026)",
        "https://careers.northwind.example/job/123",
        "21 searches, 17 pages opened, 203 seconds of research.", "Generated 5 Oct 2026",
    ):
        assert expected in text, f"missing from the Word document: {expected!r}"


def test_questions_are_a_numbered_list_and_their_notes_follow_under_bold_labels():
    data = _render(make_doc())
    assert _styles(data, NUMBERED) == [
        "Who builds a new Virtual WAN hub today, and how long does one take?",
        "Which network services stay shared after the lubricants split?",
    ]
    notes = _styles(data, CONTINUED)
    assert notes[:3] == [
        "Angle 1: Multi-cloud",
        "Listen for: hand-built hubs, weeks of lead time",
        "Alkira angle: A new region is a design change deployed in a day.",
    ]
    document = Document(io.BytesIO(data))
    listen = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("Listen for:"))
    assert listen.runs[0].text == "Listen for:" and listen.runs[0].bold
    assert not listen.runs[-1].bold


def test_lists_are_real_bullets():
    bullets = _styles(_render(make_doc()), BULLET)
    assert "Who owns the WAN contract." in bullets
    assert "A dated SD-WAN or MPLS renewal." in bullets
    assert "Clouds: Azure is the primary cloud. [1]" in bullets
    assert any(bullet.startswith("A network engineer posting lists ExpressRoute") for bullet in bullets)


def test_the_document_carries_its_title_author_and_language():
    document = Document(io.BytesIO(_render(make_doc(language="es"))))
    properties = document.core_properties
    assert properties.title == "Northwind Energy: Informe de oportunidad Alkira"
    assert properties.author == "Alkira" and properties.last_modified_by == "Alkira"
    assert properties.language == "es-ES"
    assert properties.comments == ""
    assert properties.created == WHEN.astimezone(timezone.utc)  # the local time it was made, as UTC


def test_word_does_not_open_the_file_in_compatibility_mode():
    settings = readers.docx_xml(_render(make_doc()))["word/settings.xml"]
    assert 'w:name="compatibilityMode"' in settings
    assert 'w:name="compatibilityMode" w:uri="http://schemas.microsoft.com/office/word" w:val="15"' in settings


def test_a_very_long_company_name_still_fits_the_title_property():
    company = {**SAMPLE_DOC["company"], "name": "Northwind " * 60}
    title = Document(io.BytesIO(_render(make_doc(company=company)))).core_properties.title
    assert 0 < len(title) <= 255


def test_reference_addresses_and_the_company_site_are_real_hyperlinks():
    data = _render(make_doc())
    assert readers.docx_links(data) == [
        "https://www.northwind.example",
        "https://careers.northwind.example/job/123",
        "https://www.northwind.example/annual-report.pdf",
    ]
    assert readers.docx_external_targets(data) == [(readers.HYPERLINK, link) for link in readers.docx_links(data)]


@pytest.mark.parametrize("address", [
    "javascript:alert(document.domain)", "file:///etc/passwd", "ftp://example.com/x", "//example.com/x",
    "http://127.0.0.1/admin", "https://internal.corp/wiki", "mailto:someone@example.com", "#top",
    "\\\\fileserver\\share\\brief.docx", "https://user:secret@example.com\x00/",
])
def test_an_address_that_is_not_a_public_web_page_is_never_a_link(address):
    doc = make_doc()
    doc["references"][0]["url"] = address
    doc["company"]["website"] = address
    data = _render(doc)
    assert readers.docx_links(data) == ["https://www.northwind.example/annual-report.pdf"]
    assert len(readers.docx_external_targets(data)) == 1
    assert address.replace("\x00", "") in readers.docx_text(data)  # still shown, as plain text


def test_a_link_shows_the_address_it_opens_never_one_dressed_up_as_another():
    doc = make_doc()
    doc["references"][0]["url"] = "https://paypal.example@evil.example/login"
    doc["references"][1]["url"] = "https://es.wikipedia.example/wiki/México"
    document = Document(io.BytesIO(_render(doc)))
    links = [link for paragraph in document.paragraphs for link in paragraph.hyperlinks][1:]
    assert [(link.text, link.address) for link in links] == [
        ("https://evil.example/login", "https://evil.example/login"),
        ("https://es.wikipedia.example/wiki/México", "https://es.wikipedia.example/wiki/M%C3%A9xico"),
    ]
    assert "paypal.example" not in readers.docx_text(_render(doc))


def test_the_document_is_plain_no_fields_tables_text_boxes_columns_or_macros():
    data = _render(make_doc())
    parts = readers.docx_xml(data)
    body = parts["word/document.xml"]
    for banned in ("w:fldChar", "w:instrText", "w:fldSimple", "<w:tbl>", "w:txbxContent", "w:drawing", "w:pict",
                   "mc:AlternateContent", "w:object"):
        assert banned not in body, banned
    assert 'w:num="2"' not in body  # one column
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        names = package.namelist()
        assert package.testzip() is None
    assert not any("vbaProject" in name or name.endswith(".bin") for name in names)
    assert "macroEnabled" not in parts["[Content_Types].xml"]
    assert "docProps/thumbnail.jpeg" not in names  # the template's picture of an empty page


def test_text_written_to_look_like_markup_or_a_field_stays_text():
    doc = make_doc()
    hostile = '</w:t></w:r><w:fldSimple w:instr="HYPERLINK http://evil.example"/> { INCLUDETEXT "C:\\\\secret" } =cmd|calc'
    doc["unconfirmed"] = [hostile]
    doc["angles"][0]["title"] = '<script>alert(1)</script> & "quotes"'
    data = _render(doc)
    assert hostile in _styles(data, BULLET)
    assert ("Heading 2", 'Angle 1: <script>alert(1)</script> & "quotes"') in readers.docx_headings(data)
    assert "<w:fldSimple" not in readers.docx_xml(data)["word/document.xml"]  # only ever as escaped text
    assert len(readers.docx_external_targets(data)) == 3


def test_characters_xml_cannot_hold_are_dropped_and_every_other_script_is_kept():
    company = {**SAMPLE_DOC["company"], "name": "Anker Innovations 安克创新",
               "identity_note": "Known\x00 as\x0b Anker 🚀"}
    people = [{"name": "龚银", "role": "CIO", "note": "Quoted in the case\ufffe study", "sources": [1]}]
    data = _render(make_doc(company=company, people=people))
    text = readers.docx_text(data)
    assert ("Title", "Anker Innovations 安克创新") in readers.docx_headings(data)
    assert "Known as Anker 🚀" in text and "龚银, CIO: Quoted in the case study [1]" in text


def test_a_brief_with_no_angles_leaves_the_section_out():
    fit = {"score": 1, "verdict": "No use case found.", "lead": ""}
    data = _render(make_doc(angles=[], fit=fit, people=[], questions=[], unconfirmed=[], raise_score=[], references=[]))
    assert [text for _, text in readers.docx_headings(data)] == ["Northwind Energy", "Alkira Fit", "For the engineer"]
    text = readers.docx_text(data)
    assert "1 / 5" in text and "Weak fit" in text and "Lead with" not in text
    assert readers.docx_links(data) == ["https://www.northwind.example"]


def test_one_angle_and_three_angles_each_get_their_heading():
    one = readers.docx_headings(_render(make_doc(angles=make_doc()["angles"][:1])))
    assert [text for style, text in one if style == "Heading 2"] == ["Angle 1: Hand-built Azure network"]
    third = {
        **copy.deepcopy(SAMPLE_DOC["angles"][0]), "title": "Firewalls copied per hub",
        "story": {"id": "metric", "customer": "", "result": "Firewall reduction: 73% (up to 82% in some accounts)."},
    }
    data = _render(make_doc(angles=[*copy.deepcopy(SAMPLE_DOC["angles"]), third]))
    assert [text for style, text in readers.docx_headings(data) if style == "Heading 2"][-1] == (
        "Angle 3: Firewalls copied per hub"
    )
    text = readers.docx_text(data)
    assert "Proof point: Firewall reduction: 73% (up to 82% in some accounts)." in text
    assert "No customer story matched this angle." in text


def test_a_brief_with_its_optional_fields_empty_still_renders():
    doc = make_doc(
        company={"name": "", "legal_name": "", "ticker": "", "website": "", "identity_note": ""},
        stats={key: "" for key in SAMPLE_DOC["stats"]},
        snapshot={key: {"text": "", "sources": []} for key in SAMPLE_DOC["snapshot"]},
        people=[{"name": "", "role": "", "note": "", "sources": []}],
    )
    doc["angles"][0].update(title="", alkira="", evidence=[], story={"id": "", "customer": "", "result": ""})
    doc["questions"][0].update(listen_for="", alkira_angle="", angle=0)
    doc["references"][0].update(title="", url="", date="")
    data = _render(doc)
    headings = readers.docx_headings(data)
    assert headings[0] == ("Title", "Alkira opportunity brief")
    assert ("Heading 2", "Angle 1") in headings
    text = readers.docx_text(data)
    assert text.count("Not found in public sources.") == 6 and "source undated" in text
    assert not any(not paragraph.strip() for _, paragraph in readers.docx_paragraphs(data))  # no empty paragraphs


def test_a_document_stored_by_older_code_renders():
    old = make_doc()
    del old["version"]
    for angle in old["angles"]:
        for key in ("deal_date", "deal_status", "deal_pending_quote"):
            del angle[key]
    for question in old["questions"]:
        del question["angle"]
    for reference in old["references"]:
        del reference["source_type"], reference["open_posting"]
    doc = brief_doc.load(json.dumps(old))
    text = readers.docx_text(_render(doc))
    assert "Lubricants separation" in text and "Pending deal" not in text and "Second-hand" in text


def test_a_spanish_brief_is_labelled_in_spanish():
    data = _render(make_doc(language="es"))
    assert [text for style, text in readers.docx_headings(data) if style == "Heading 1"] == [
        "Ajuste Alkira", "Por qué esta cuenta, por qué ahora", "Pregunte esto", "Para el ingeniero",
        "Con quién hablar", "Lo que no pudimos confirmar", "Qué subiría la puntuación", "Referencias",
    ]
    text = readers.docx_text(data)
    for expected in ("Ángulo 1: Hand-built Azure network", "Caso de uso: Multinube", "Qué escuchar:",
                     "Operación pendiente, anunciada el 20 feb 2026", "(Fuente directa; fuente con fecha 23 sep 2026)",
                     "Generado el 5 oct 2026", "CONFIDENCIAL"):
        assert expected in text or expected in " ".join(readers.docx_xml(data).values()), expected


def test_the_language_asked_for_wins_over_the_language_of_the_document():
    assert "Pregunte esto" in readers.docx_text(_render(make_doc(), "es"))


def test_rendering_does_not_change_the_document_it_is_given():
    doc = make_doc()
    before = copy.deepcopy(doc)
    _render(doc)
    assert doc == before
