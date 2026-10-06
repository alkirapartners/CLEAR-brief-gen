"""Read an exported brief back: the text and links of a PDF, the paragraphs and links of a Word file."""

import io
import re
import zipfile

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from pypdf import PdfReader

_WHITESPACE = re.compile(r"\s+")


def squash(text: str) -> str:
    """Text with every run of whitespace as one space, so a wrapped line reads as it was written."""
    return _WHITESPACE.sub(" ", text).strip()


# ── PDF ──────────────────────────────────────────────────────────────────────

def pdf_pages(data: bytes) -> list[str]:
    """The text of each page, in reading order."""
    return [squash(page.extract_text()) for page in PdfReader(io.BytesIO(data)).pages]


def pdf_text(data: bytes) -> str:
    return " ".join(pdf_pages(data))


def _annotations(data: bytes) -> list[dict]:
    reader = PdfReader(io.BytesIO(data))
    return [note.get_object() for page in reader.pages for note in page.get("/Annots") or []]


def pdf_web_links(data: bytes) -> list[str]:
    """Every address a link in the PDF opens."""
    actions = [note.get("/A") for note in _annotations(data)]
    return [str(action["/URI"]) for action in actions if action is not None and "/URI" in action]


def pdf_jumps(data: bytes) -> int:
    """How many links jump to another place in the PDF."""
    return sum(1 for note in _annotations(data) if "/Dest" in note)


def pdf_fonts_used(data: bytes) -> set[str]:
    reader = PdfReader(io.BytesIO(data))
    names: set[str] = set()
    for page in reader.pages:
        for font in page["/Resources"]["/Font"].values():
            names.add(str(font.get_object()["/BaseFont"]))
    return names


def pdf_page_size(data: bytes) -> tuple[float, float]:
    box = PdfReader(io.BytesIO(data)).pages[0].mediabox
    return float(box.width), float(box.height)


def pdf_metadata(data: bytes) -> dict[str, str]:
    return {key: str(value) for key, value in (PdfReader(io.BytesIO(data)).metadata or {}).items()}


# ── Word ─────────────────────────────────────────────────────────────────────

def docx_paragraphs(data: bytes) -> list[tuple[str, str]]:
    """Each paragraph's style name and its text, links included."""
    document = Document(io.BytesIO(data))
    return [(paragraph.style.name, paragraph.text) for paragraph in document.paragraphs]


def docx_text(data: bytes) -> str:
    return squash(" ".join(text for _, text in docx_paragraphs(data)))


def docx_headings(data: bytes) -> list[tuple[str, str]]:
    return [(style, text) for style, text in docx_paragraphs(data) if style.startswith(("Heading", "Title"))]


def docx_links(data: bytes) -> list[str]:
    """Every address a hyperlink in the document opens."""
    document = Document(io.BytesIO(data))
    return [link.address for paragraph in document.paragraphs for link in paragraph.hyperlinks]


def docx_external_targets(data: bytes) -> list[tuple[str, str]]:
    """Everything outside the file that any part of it points at: relationship type and target."""
    document = Document(io.BytesIO(data))
    found: list[tuple[str, str]] = []
    for part in document.part.package.iter_parts():
        found += [(rel.reltype, rel.target_ref) for rel in part.rels.values() if rel.is_external]
    return found


def docx_xml(data: bytes) -> dict[str, str]:
    """The XML of every part of the package, by its name."""
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        return {
            name: package.read(name).decode("utf-8")
            for name in package.namelist() if name.endswith((".xml", ".rels"))
        }


HYPERLINK = RELATIONSHIP_TYPE.HYPERLINK
