"""Word export of a brief stored as a JSON document.

Deliberately plain, like Markdown made into a document: one column, real
Word headings, real bullet and numbered lists, bold run-in labels, and the
sources as real hyperlinks. It says what the PDF says, in the same order
(export_content.py), and is meant to be edited or pasted into an email.

Everything in a brief was written by a model from other people's pages, so
every string goes in as text and nothing else: no field codes, no macros,
and the only things the file points at outside itself are hyperlinks to
public http(s) pages (export_content.link_target). Legacy markdown briefs
have no Word export.
"""

import io
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Callable

from docx import Document
from docx.document import Document as WordDocument
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run

import brief_compat
import export_content
import i18n
from brief_doc import SNAPSHOT_KEYS, Angle, BriefDoc, Person, Question, Reference

Labels = dict[str, str]

AUTHOR = "Alkira"
# A font every Word, Pages and Google Docs has, so none of them substitutes one or warns.
FONT = "Arial"
INK = RGBColor(0x14, 0x14, 0x14)
MUTED = RGBColor(0x5F, 0x66, 0x73)
LINK = RGBColor(0x2D, 0x58, 0xF2)
BODY_SIZE = Pt(10.5)
HEADING_SIZES: dict[str, Pt] = {"Title": Pt(24), "Heading 1": Pt(15), "Heading 2": Pt(12)}
MARGIN = Inches(1)
LETTER_WIDTH, LETTER_HEIGHT = Inches(8.5), Inches(11)
# Word refuses a document property longer than this.
MAX_PROPERTY_CHARS = 255
# What each brief language is called in a document, for spelling and hyphenation.
LANGUAGE_TAGS: dict[str, str] = {"en": "en-US", "es": "es-ES"}
SEPARATOR = " · "
BULLET, NUMBERED, CONTINUED, QUOTE = "List Bullet", "List Number", "List Continue", "Quote"
THUMBNAIL = "http://schemas.openxmlformats.org/package/2006/relationships/metadata/thumbnail"
# The file format of Word 2013 and later. The template python-docx starts from is marked as Word 2010's.
WORD_FORMAT = "15"
# Characters XML 1.0 cannot hold: control codes, lone surrogates and the two non-characters.
_NOT_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


@dataclass(frozen=True)
class _Context:
    """What every section is written from: the Word document being built, the brief and its language."""

    document: WordDocument
    doc: BriefDoc
    labels: Labels
    language: str


def clean(text: str) -> str:
    """Text as one line that XML can hold, whatever was written."""
    return " ".join(_NOT_XML.sub("", text).split())


# ── Writing ──────────────────────────────────────────────────────────────────

def _write(paragraph: Paragraph, text: str, bold: bool = False, muted: bool = False) -> None:
    """Add text to a paragraph as text and nothing else. Empty text adds nothing."""
    content = clean(text)
    if not content:
        return
    # A run that follows another keeps the space between them, which clean() would trim.
    lead = " " if text[:1].isspace() and paragraph.runs else ""
    tail = " " if text[-1:].isspace() else ""
    run = paragraph.add_run(f"{lead}{content}{tail}")
    run.bold = bold or None
    if muted:
        run.font.color.rgb = MUTED


def _labelled(ctx: _Context, label: str, text: str, style: str | None = None) -> None:
    """A paragraph that opens with a bold label: "Listen for: ...". Skipped when there is nothing to say."""
    if not clean(text):
        return
    paragraph = ctx.document.add_paragraph(style=style)
    _write(paragraph, f"{label}:", bold=True)
    _write(paragraph, f" {text}")


def _line(ctx: _Context, text: str, style: str | None = None, muted: bool = False) -> None:
    """A paragraph of plain text. Skipped when the text is empty."""
    if clean(text):
        _write(ctx.document.add_paragraph(style=style), text, muted=muted)


def _link(paragraph: Paragraph, text: str, url: str) -> None:
    """Text that opens a web page. ``url`` has already been through export_content.link_target."""
    relation = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation)
    element = OxmlElement("w:r")
    hyperlink.append(element)
    paragraph._p.append(hyperlink)
    run = Run(element, paragraph)
    run.text = clean(text)
    run.font.color.rgb = LINK
    run.font.underline = True


def _address(paragraph: Paragraph, shown: str, url: str) -> None:
    """An address: a link when it is a public web page, plain text when it is anything else."""
    target = export_content.link_target(url)
    if target is None:
        _write(paragraph, shown)
    elif clean(shown):
        _link(paragraph, shown, target)


# ── Sections ─────────────────────────────────────────────────────────────────

def _opening(ctx: _Context) -> None:
    """Whose brief this is: the company, what the document is, and the company's basics."""
    doc, labels = ctx.doc, ctx.labels
    ctx.document.add_heading(clean(doc["company"]["name"]) or labels["brief_title"], level=0)
    researched = i18n.readable_date(doc["generated"], ctx.language)
    when = labels["researched_on"].format(date=researched) if researched else ""
    _line(ctx, SEPARATOR.join(part for part in (labels["brief_title"], when) if part), muted=True)
    identity = export_content.identity(doc, labels)
    if identity.facts or identity.site_url:
        paragraph = ctx.document.add_paragraph()
        _write(paragraph, SEPARATOR.join(identity.facts))
        if identity.site_url:
            _write(paragraph, SEPARATOR if identity.facts else "")
            _link(paragraph, identity.site, identity.site_url)
    _labelled(ctx, labels["identity"], identity.note)
    stats = export_content.stat_pills_for(doc, labels)
    if stats:
        paragraph = ctx.document.add_paragraph()
        for index, (name, value) in enumerate(stats):
            _write(paragraph, f"{SEPARATOR if index else ''}{name}:", bold=True)
            _write(paragraph, f" {value}")


def _fit(ctx: _Context) -> None:
    fit = ctx.doc["fit"]
    score = f"{export_content.score_text(fit['score'])}{SEPARATOR}{export_content.fit_label(fit['score'], ctx.labels)}"
    _write(ctx.document.add_paragraph(), score, bold=True)
    _line(ctx, fit["verdict"])
    _labelled(ctx, ctx.labels["lead"], fit["lead"])
    _labelled(ctx, ctx.labels["stat_cloud_network"], ctx.doc["stats"]["cloud_network"])


def _proof(ctx: _Context, angle: Angle) -> None:
    """The angle's proof. A knowledge-base figure says that it is not a customer's result."""
    proof = export_content.proof(angle["story"], ctx.labels)
    if proof is None:
        return
    said = brief_compat.proof_text(angle["story"])
    _labelled(ctx, proof.label, f"{said} {proof.note}" if proof.note else said)


def _angle(ctx: _Context, number: int, angle: Angle) -> None:
    labels = ctx.labels
    title = clean(angle["title"])
    numbered = f"{labels['angle']} {number}"
    ctx.document.add_heading(f"{numbered}: {title}" if title else numbered, level=2)
    deal = export_content.deal_text(angle, labels, ctx.language)
    use_case = export_content.use_case_name(angle["use_case"], labels)
    _labelled(ctx, labels["use_case"], SEPARATOR.join(part for part in (use_case, deal) if part))
    quote = clean(angle["deal_pending_quote"])
    _line(ctx, f"“{quote}”" if quote else "", style=QUOTE)
    _labelled(ctx, labels["alkira_answer"], angle["alkira"])
    facts = [line for line in angle["evidence"] if clean(line["text"])]
    if facts:
        _write(ctx.document.add_paragraph(), f"{labels['evidence']}:", bold=True)
    for line in facts:
        date = export_content.evidence_date(line, ctx.doc["references"], labels, ctx.language)
        note = f" ({date.note})" if date.note else ""
        _line(ctx, f"{line['text']}{note}{brief_compat.cite(line['sources'])}", style=BULLET)
    _proof(ctx, angle)


def _angles(ctx: _Context) -> None:
    for number, angle in enumerate(ctx.doc["angles"], start=1):
        _angle(ctx, number, angle)


def _question(ctx: _Context, item: Question) -> None:
    """A numbered question, then which angle it is about and its two notes, indented under it."""
    _line(ctx, item["question"], style=NUMBERED)
    number, angles = item["angle"], ctx.doc["angles"]
    if 0 < number <= len(angles):
        use_case = export_content.use_case_name(angles[number - 1]["use_case"], ctx.labels)
        _labelled(ctx, f"{ctx.labels['angle']} {number}", use_case, style=CONTINUED)
    _labelled(ctx, ctx.labels["listen_for"], item["listen_for"], style=CONTINUED)
    _labelled(ctx, ctx.labels["alkira_angle"], item["alkira_angle"], style=CONTINUED)


def _questions(ctx: _Context) -> None:
    for item in ctx.doc["questions"]:
        if clean(item["question"]):
            _question(ctx, item)


def _engineer_sheet(ctx: _Context) -> None:
    """Every line of the technical snapshot. One the research did not find says so."""
    for key in SNAPSHOT_KEYS:
        line = ctx.doc["snapshot"][key]
        found = f"{line['text']}{brief_compat.cite(line['sources'])}" if clean(line["text"]) else ""
        _labelled(ctx, ctx.labels[f"snap_{key}"], found or ctx.labels["not_found_public"], style=BULLET)


def _person(ctx: _Context, person: Person) -> None:
    name, role, note = clean(person["name"]), clean(person["role"]), clean(person["note"])
    if not name and not role:
        return
    paragraph = ctx.document.add_paragraph(style=BULLET)
    _write(paragraph, name or role, bold=True)
    _write(paragraph, f", {role}" if name and role else "")
    _write(paragraph, f": {note}" if note else "")
    _write(paragraph, brief_compat.cite(person["sources"]))


def _people(ctx: _Context) -> None:
    for person in ctx.doc["people"]:
        _person(ctx, person)


def _bullets(key: str) -> Callable[[_Context], None]:
    def write(ctx: _Context) -> None:
        for item in ctx.doc[key]:
            _line(ctx, item, style=BULLET)

    return write


def _reference(ctx: _Context, reference: Reference) -> None:
    """One source: its number and title, what kind of source it is and how it is dated, then its address."""
    labels = ctx.labels
    kind = export_content.source_type_label(reference, labels)
    dated = export_content.reference_date(reference, labels, ctx.language).note
    paragraph = ctx.document.add_paragraph()
    _write(paragraph, f"[{reference['n']}]", bold=True)
    _write(paragraph, f" {export_content.reference_title(reference, labels)}")
    _write(paragraph, f" ({kind}; {dated})")
    if clean(reference["url"]):
        paragraph.add_run().add_break()
        _address(paragraph, reference["url"], reference["url"])


def _references(ctx: _Context) -> None:
    for reference in ctx.doc["references"]:
        _reference(ctx, reference)


def _closing(ctx: _Context, generated_at: datetime) -> None:
    """How much research stands behind the brief, and when this file was made."""
    made = i18n.readable_date(generated_at.strftime("%Y-%m-%d"), ctx.language)
    note = f"{export_content.research_note(ctx.doc, ctx.labels)}. {ctx.labels['generated_on'].format(date=made)}."
    _line(ctx, note, muted=True)


# How each section of export_content.outline is written.
WRITERS: dict[str, Callable[[_Context], None]] = {
    "fit": _fit,
    "why_now": _angles,
    "ask_this": _questions,
    "for_engineer": _engineer_sheet,
    "who_to_talk_to": _people,
    "unconfirmed": _bullets("unconfirmed"),
    "raise_score": _bullets("raise_score"),
    "references": _references,
}


# ── The document itself ──────────────────────────────────────────────────────

def _plain_fonts(document: WordDocument) -> None:
    """One common font in dark ink throughout, in place of the template's theme fonts and blue headings."""
    defaults = document.styles.element.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}")
    fonts = defaults.find(qn("w:rFonts"))
    fonts.attrib.clear()
    for script in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(script), FONT)
    document.styles["Normal"].font.size = BODY_SIZE
    for name, size in HEADING_SIZES.items():
        style = document.styles[name]
        properties = style.element.get_or_add_rPr()
        for theme_font in properties.findall(qn("w:rFonts")):
            properties.remove(theme_font)
        style.font.color.rgb = INK
        style.font.size = size
        style.font.bold = True
    # The template rules off its title in blue. A plain document has no rules.
    title = document.styles["Title"].element.pPr
    for border in title.findall(qn("w:pBdr")):
        title.remove(border)


def _set_language(document: WordDocument, tag: str) -> None:
    """Tell Word which language to check the spelling in."""
    language = document.styles.element.find(
        f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}/{qn('w:lang')}"
    )
    language.set(qn("w:val"), tag)
    document.core_properties.language = tag


def _current_format(document: WordDocument) -> None:
    """Mark the file as a current Word document, so Word does not open it in "Compatibility Mode"."""
    for setting in document.settings.element.iter(qn("w:compatSetting")):
        if setting.get(qn("w:name")) == "compatibilityMode":
            setting.set(qn("w:val"), WORD_FORMAT)


def _drop_thumbnail(document: WordDocument) -> None:
    """The template carries a picture of an empty page as its preview. Leave it out."""
    relations = document.part.package.rels
    for key in [key for key, relation in relations.items() if relation.reltype == THUMBNAIL]:
        del relations[key]


def _new_document(title: str, generated_at: datetime, language: str, confidential: str) -> WordDocument:
    """An empty US Letter document with its properties set and the confidentiality mark at the head of each page."""
    document = Document()
    _plain_fonts(document)
    _set_language(document, LANGUAGE_TAGS.get(language, LANGUAGE_TAGS[i18n.DEFAULT_LANGUAGE]))
    _current_format(document)
    _drop_thumbnail(document)
    properties = document.core_properties
    properties.title = clean(title)[:MAX_PROPERTY_CHARS]
    properties.author = properties.last_modified_by = AUTHOR
    properties.comments = ""
    # Word reads these as UTC. A time with no zone is the server's local time.
    properties.created = properties.modified = generated_at.astimezone(timezone.utc).replace(tzinfo=None)
    section = document.sections[0]
    section.page_width, section.page_height = LETTER_WIDTH, LETTER_HEIGHT
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = MARGIN
    _write(section.header.paragraphs[0], confidential.upper(), muted=True)
    return document


def render(doc: BriefDoc, generated_at: datetime, language: str | None = None) -> bytes:
    """The brief document as a Word file, in the document's own language by default."""
    code = i18n.normalize(language or doc["language"])
    labels = i18n.labels(code)
    title = export_content.document_title(doc, labels)
    ctx = _Context(_new_document(title, generated_at, code, labels["confidential"]), doc, labels, code)
    _opening(ctx)
    for section in export_content.outline(doc, labels):
        ctx.document.add_heading(section.heading, level=1)
        WRITERS[section.key](ctx)
    _closing(ctx, generated_at)
    buffer = io.BytesIO()
    ctx.document.save(buffer)
    return buffer.getvalue()
